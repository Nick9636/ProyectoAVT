# Adriana Nicole Guzman Ahuatzi
#01/04/2026
# Rutas para reportes estadísticos del sistema de afiliación de jugadores a ligas deportivas.
from flask import Blueprint, render_template, request, jsonify, session, redirect, url_for, send_file
from extensiones import mysql
import io

reportes = Blueprint('reportes', __name__, url_prefix="/reportes")


def login_requerido(f):
    from functools import wraps
    @wraps(f)
    def decorador(*args, **kwargs):
        if 'id_usuario' not in session:
            return redirect(url_for('auth.iniciarSesion'))
        return f(*args, **kwargs)
    return decorador


# ── General ───────────────────────────────────────
@reportes.route("/")
@login_requerido
def reporteGeneral():
    cur = mysql.connection.cursor()
    cur.execute("SELECT COUNT(*) AS t FROM jugador")
    total_jugadores = cur.fetchone()['t']
    cur.execute("SELECT COUNT(*) AS t FROM liga")
    total_ligas = cur.fetchone()['t']
    cur.execute("SELECT COUNT(*) AS t FROM equipo")
    total_equipos = cur.fetchone()['t']
    cur.execute("SELECT COUNT(*) AS t FROM pago WHERE estatus='Completado'")
    pagos_completados = cur.fetchone()['t']
    cur.close()
    return render_template("reportes/reporteGeneral.html",
        total_jugadores=total_jugadores,
        total_ligas=total_ligas,
        total_equipos=total_equipos,
        pagos_completados=pagos_completados)


# ── Ingresos ──────────────────────────────────────
@reportes.route("/ingresos")
@login_requerido
def reporteIngresos():
    fecha_desde = request.args.get('fecha_desde', '')
    fecha_hasta = request.args.get('fecha_hasta', '')
    estatus     = request.args.get('estatus', '')

    cur = mysql.connection.cursor()

    condiciones = []
    valores     = []
    if fecha_desde:
        condiciones.append("p.fecha_pago >= %s"); valores.append(fecha_desde)
    if fecha_hasta:
        condiciones.append("p.fecha_pago <= %s"); valores.append(fecha_hasta)
    if estatus:
        condiciones.append("p.estatus = %s"); valores.append(estatus)

    where = f"WHERE {' AND '.join(condiciones)}" if condiciones else ""

    cur.execute(f"""
        SELECT p.id_pago, p.fecha_pago, p.estatus, p.metodo_pago, p.referencia,
               CONCAT(j.apellido_paterno,' ',j.apellido_materno,' ',j.nombres) AS nombre_jugador,
               a.numero_registro
        FROM pago p
        JOIN afiliacion a ON a.id_afiliacion = p.id_afiliacion
        JOIN jugador j    ON j.id_jugador    = a.id_jugador
        {where}
        ORDER BY p.fecha_pago DESC
    """, valores)
    pagos = cur.fetchall()

    cur.execute("SELECT COUNT(*) AS t FROM pago WHERE estatus='Completado'")
    completados = cur.fetchone()['t']
    cur.execute("SELECT COUNT(*) AS t FROM pago WHERE estatus='Pendiente'")
    pendientes = cur.fetchone()['t']
    cur.execute("SELECT COUNT(*) AS t FROM pago WHERE estatus='Cancelado'")
    cancelados = cur.fetchone()['t']
    cur.close()

    return render_template("reportes/reporteIngresos.html",
        pagos=pagos,
        completados=completados,
        pendientes=pendientes,
        cancelados=cancelados,
        filtro_desde=fecha_desde,
        filtro_hasta=fecha_hasta,
        filtro_estatus=estatus)


@reportes.route("/ingresos/excel")
@login_requerido
def exportarIngresosExcel():
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return jsonify({'ok': False, 'mensaje': 'openpyxl no instalado'}), 500

    cur = mysql.connection.cursor()
    cur.execute("""
        SELECT p.id_pago, p.fecha_pago, p.estatus, p.metodo_pago, p.referencia,
               CONCAT(j.apellido_paterno,' ',j.apellido_materno,' ',j.nombres) AS nombre_jugador,
               a.numero_registro
        FROM pago p
        JOIN afiliacion a ON a.id_afiliacion = p.id_afiliacion
        JOIN jugador j    ON j.id_jugador    = a.id_jugador
        ORDER BY p.fecha_pago DESC
    """)
    pagos = cur.fetchall()
    cur.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ingresos"

    header_fill = PatternFill("solid", fgColor="7A0C0C")
    header_font = Font(bold=True, color="FFFFFF")
    headers = ["ID", "Fecha", "Jugador", "No. Registro", "Método", "Referencia", "Estatus"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row, p in enumerate(pagos, 2):
        ws.cell(row=row, column=1, value=p['id_pago'])
        ws.cell(row=row, column=2, value=str(p['fecha_pago']))
        ws.cell(row=row, column=3, value=p['nombre_jugador'])
        ws.cell(row=row, column=4, value=p['numero_registro'] or '—')
        ws.cell(row=row, column=5, value=p['metodo_pago'] or '—')
        ws.cell(row=row, column=6, value=p['referencia'] or '—')
        ws.cell(row=row, column=7, value=p['estatus'])

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True, download_name='reporte_ingresos.xlsx')


# ── Inscripción ───────────────────────────────────
@reportes.route("/inscripcion")
@login_requerido
def reporteInscripcion():
    cur = mysql.connection.cursor()

    cur.execute("""
        SELECT j.id_jugador,
               CONCAT(j.apellido_paterno,' ',j.apellido_materno,' ',j.nombres) AS nombre,
               j.curp, j.correo_electronico,
               e.estatus AS estatus_expediente,
               e.fecha_creacion
        FROM jugador j
        LEFT JOIN expediente e ON e.id_jugador = j.id_jugador
        ORDER BY e.fecha_creacion DESC
    """)
    jugadores = cur.fetchall()

    cur.execute("SELECT COUNT(*) AS t FROM jugador")
    total = cur.fetchone()['t']
    cur.execute("SELECT COUNT(*) AS t FROM expediente WHERE estatus='activo'")
    activos = cur.fetchone()['t']
    cur.execute("""
        SELECT COUNT(*) AS t FROM expediente
        WHERE DATE(fecha_creacion) >= DATE_SUB(CURDATE(), INTERVAL 30 DAY)
    """)
    nuevos_mes = cur.fetchone()['t']
    cur.close()

    return render_template("reportes/reporteInscripcion.html",
        jugadores=jugadores,
        total=total,
        activos=activos,
        nuevos_mes=nuevos_mes)


@reportes.route("/inscripcion/excel")
@login_requerido
def exportarInscripcionExcel():
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return jsonify({'ok': False, 'mensaje': 'openpyxl no instalado'}), 500

    cur = mysql.connection.cursor()
    cur.execute("""
        SELECT j.id_jugador,
               CONCAT(j.apellido_paterno,' ',j.apellido_materno,' ',j.nombres) AS nombre,
               j.curp, j.correo_electronico,
               e.estatus AS estatus_expediente,
               e.fecha_creacion
        FROM jugador j
        LEFT JOIN expediente e ON e.id_jugador = j.id_jugador
        ORDER BY e.fecha_creacion DESC
    """)
    jugadores = cur.fetchall()
    cur.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inscripciones"

    header_fill = PatternFill("solid", fgColor="7A0C0C")
    header_font = Font(bold=True, color="FFFFFF")
    headers = ["ID", "Nombre", "CURP", "Correo", "Estatus", "Fecha Registro"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row, j in enumerate(jugadores, 2):
        ws.cell(row=row, column=1, value=j['id_jugador'])
        ws.cell(row=row, column=2, value=j['nombre'])
        ws.cell(row=row, column=3, value=j['curp'])
        ws.cell(row=row, column=4, value=j['correo_electronico'] or '—')
        ws.cell(row=row, column=5, value=j['estatus_expediente'] or '—')
        ws.cell(row=row, column=6, value=str(j['fecha_creacion']) if j['fecha_creacion'] else '—')

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 20

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True, download_name='reporte_inscripciones.xlsx')


# ── Sistema ───────────────────────────────────────
@reportes.route("/sistema")
@login_requerido
def reporteSistema():
    cur = mysql.connection.cursor()

    cur.execute("SELECT COUNT(*) AS t FROM usuario")
    total_usuarios = cur.fetchone()['t']
    cur.execute("SELECT COUNT(*) AS t FROM liga")
    total_ligas = cur.fetchone()['t']
    cur.execute("SELECT COUNT(*) AS t FROM equipo")
    total_equipos = cur.fetchone()['t']
    cur.execute("SELECT COUNT(*) AS t FROM jugador")
    total_jugadores = cur.fetchone()['t']
    cur.execute("SELECT COUNT(*) AS t FROM afiliacion")
    total_afiliaciones = cur.fetchone()['t']
    cur.execute("SELECT COUNT(*) AS t FROM pago WHERE estatus='Completado'")
    pagos_ok = cur.fetchone()['t']
    cur.execute("SELECT COUNT(*) AS t FROM pago WHERE estatus='Pendiente'")
    pagos_pend = cur.fetchone()['t']
    cur.execute("SELECT COUNT(*) AS t FROM pago WHERE estatus='Cancelado'")
    pagos_cancel = cur.fetchone()['t']
    cur.execute("""
        SELECT COUNT(*) AS t FROM autorizacion_pendiente
        WHERE estatus='Pendiente'
    """)
    solicitudes_pend = cur.fetchone()['t']
    cur.execute("""
        SELECT COUNT(*) AS t FROM autorizacion_pendiente
        WHERE estatus='Autorizado' AND DATE(fecha_solicitud) = CURDATE()
    """)
    autorizados_hoy = cur.fetchone()['t']
    cur.execute("SELECT COUNT(*) AS t FROM formato_digital")
    total_docs = cur.fetchone()['t']

    # Ligas con sus equipos y jugadores
    cur.execute("""
        SELECT l.nombre_liga, l.categoria,
               COUNT(DISTINCT e.id_equipo) AS equipos,
               COUNT(DISTINCT jl.id_jugador) AS jugadores
        FROM liga l
        LEFT JOIN equipo e ON e.id_liga = l.id_liga
        LEFT JOIN jugador_liga jl ON jl.id_liga = l.id_liga
        GROUP BY l.id_liga
        ORDER BY jugadores DESC
        LIMIT 10
    """)
    ligas = cur.fetchall()

    # Usuarios por rol
    cur.execute("SELECT rol, COUNT(*) AS total FROM usuario GROUP BY rol")
    usuarios_rol = cur.fetchall()

    cur.close()

    total_pagos = pagos_ok + pagos_pend + pagos_cancel or 1
    return render_template("reportes/reporteSistema.html",
        total_usuarios=total_usuarios,
        total_ligas=total_ligas,
        total_equipos=total_equipos,
        total_jugadores=total_jugadores,
        total_afiliaciones=total_afiliaciones,
        pagos_ok=pagos_ok,
        pagos_pend=pagos_pend,
        pagos_cancel=pagos_cancel,
        pct_ok=round(pagos_ok/total_pagos*100),
        pct_pend=round(pagos_pend/total_pagos*100),
        pct_cancel=round(pagos_cancel/total_pagos*100),
        solicitudes_pend=solicitudes_pend,
        autorizados_hoy=autorizados_hoy,
        total_docs=total_docs,
        ligas=ligas,
        usuarios_rol=usuarios_rol)
